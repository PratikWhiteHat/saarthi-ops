"""Phase 8A reporting extensions: CVSS, remediation roadmap, MD/HTML/PDF."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from saarthi_ai.exploit_confirmation.models import Severity
from saarthi_ai.reporting import cvss, pdf
from saarthi_ai.reporting.docx_renderer import render_report_docx
from saarthi_ai.reporting.html_renderer import render_report_html
from saarthi_ai.reporting.markdown_renderer import render_report_markdown
from saarthi_ai.reporting.models import (
    EngagementMeta,
    ReportFinding,
    ReportModel,
    estimate_effort,
    severity_rating_key,
)


def _model() -> ReportModel:
    xss = ReportFinding(
        finding_id="cross-site-scripting-xss",
        number=1,
        title="Cross-Site Scripting (XSS)",
        severity=Severity.HIGH,
        affected_urls=["https://target.test/search"],
        description="Reflected XSS in the search box.",
        security_risk="Session theft.",
        recommendation="Encode output.",
        proof_of_concept="Inject <script>.",
        cvss_vector="CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:N/A:N",
        cvss_score=7.5,
        remediation_effort="Medium",
    )
    cookie = ReportFinding(
        finding_id="cookie-flag-httponly-not-set",
        number=2,
        title="Cookie Flag HTTPOnly Not Set",
        severity=Severity.LOW,
        recommendation="Set HttpOnly.",
        cvss_vector="CVSS:3.1/AV:N/AC:H/PR:N/UI:R/S:U/C:L/I:N/A:N",
        cvss_score=3.1,
        remediation_effort="Low",
    )
    return ReportModel(
        engagement=EngagementMeta(
            app_name="https://target.test", company="Acme Security"
        ),
        target="https://target.test",
        executive_summary="Two issues were identified.",
        findings=[xss, cookie],
        methodology="Evidence-driven, permission-gated testing.",
        tools_used=["httpx", "nuclei", "sqlmap"],
    )


# --- CVSS --------------------------------------------------------------------


def test_cvss_base_score_known_vectors() -> None:
    critical = {"AV": "N", "AC": "L", "PR": "N", "UI": "N", "S": "U",
                "C": "H", "I": "H", "A": "H"}
    high = {"AV": "N", "AC": "L", "PR": "N", "UI": "N", "S": "U",
            "C": "H", "I": "N", "A": "N"}
    xss = {"AV": "N", "AC": "L", "PR": "N", "UI": "R", "S": "C",
           "C": "L", "I": "L", "A": "N"}
    assert cvss.base_score(critical) == 9.8
    assert cvss.base_score(high) == 7.5
    assert cvss.base_score(xss) == 6.1


def test_cvss_severity_bands() -> None:
    assert cvss.severity_from_score(0.0) is Severity.INFO
    assert cvss.severity_from_score(3.9) is Severity.LOW
    assert cvss.severity_from_score(6.9) is Severity.MEDIUM
    assert cvss.severity_from_score(8.9) is Severity.HIGH
    assert cvss.severity_from_score(9.8) is Severity.CRITICAL


def test_cvss_classify_is_band_aligned() -> None:
    # Type template used when its band matches the assigned severity.
    sqli = cvss.classify("SQL Injection", Severity.CRITICAL)
    assert sqli.severity is Severity.CRITICAL
    assert sqli.score == 9.8

    # Same type at a lower assigned severity falls back to a band-aligned vector.
    for severity in (Severity.HIGH, Severity.MEDIUM, Severity.LOW):
        result = cvss.classify("SQL Injection", severity)
        assert result.severity is severity

    # Unknown finding titles still land in the assigned band.
    for severity in (Severity.CRITICAL, Severity.HIGH, Severity.MEDIUM,
                     Severity.LOW):
        assert cvss.classify("Some novel issue", severity).severity is severity


def test_cvss_info_has_no_vector() -> None:
    result = cvss.classify("Informational note", Severity.INFO)
    assert result.available is False
    assert result.score == 0.0


# --- effort / rating key -----------------------------------------------------


def test_estimate_effort_buckets() -> None:
    assert estimate_effort("Broken Access Control", Severity.HIGH) == "High"
    assert estimate_effort("Cookie without Secure flag", Severity.LOW) == "Low"
    assert estimate_effort("Some medium issue", Severity.MEDIUM) == "Medium"


def test_severity_rating_key_shape() -> None:
    key = severity_rating_key()
    assert [row["priority"] for row in key] == ["P1", "P2", "P3", "P4", "P5"]


def test_remediation_roadmap_orders_worst_first() -> None:
    roadmap = _model().remediation_roadmap
    assert [item.priority for item in roadmap] == ["P2", "P4"]
    assert roadmap[0].effort == "Medium"


# --- renderers ---------------------------------------------------------------


def test_markdown_renderer_covers_all_sections(tmp_path: Path) -> None:
    out = render_report_markdown(_model(), tmp_path / "r.md")
    text = out.read_text()
    assert "## Remediation Roadmap" in text
    assert "## Detailed Findings" in text
    assert "### Methodology" in text
    assert "### Severity Rating Key" in text
    assert "CVSS v3.1" in text
    assert "Cross-Site Scripting (XSS)" in text
    assert "httpx" in text


def test_html_renderer_is_self_contained(tmp_path: Path) -> None:
    out = render_report_html(_model(), tmp_path / "r.html")
    html = out.read_text()
    assert html.startswith("<!doctype html>")
    assert "<style>" in html  # inline CSS, no external assets
    assert "Remediation Roadmap" in html
    assert "Severity Rating Key" in html
    assert "CVSS:3.1/AV:N" in html
    assert "Cross-Site Scripting (XSS)" in html


def test_docx_fallback_has_roadmap_and_appendix(tmp_path: Path) -> None:
    out = render_report_docx(_model(), tmp_path / "r.docx", template_path=None)
    document = Document(str(out))
    full = "\n".join(p.text for p in document.paragraphs)
    full += "\n" + "\n".join(
        c.text for t in document.tables for r in t.rows for c in r.cells
    )
    assert "Remediation Roadmap" in full
    assert "Appendix" in full
    assert "Methodology" in full
    assert "Severity Rating Key" in full
    assert "CVSS v3.1" in full


# --- PDF (best-effort) -------------------------------------------------------


def test_pdf_returns_none_without_converter(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(pdf, "find_soffice", lambda: None)
    source = tmp_path / "r.docx"
    source.write_bytes(b"stub")
    assert pdf.render_report_pdf(source, tmp_path / "r.pdf") is None


def test_pdf_returns_none_when_source_missing(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setattr(pdf, "find_soffice", lambda: "/usr/bin/soffice")
    missing = tmp_path / "nope.docx"
    assert pdf.render_report_pdf(missing, tmp_path / "r.pdf") is None
