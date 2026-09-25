"""Phase 8A reporting: builder, renderers, and tracked-workflow tests."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from saarthi_ai.exploit_confirmation.models import Severity
from saarthi_ai.knowledge.models import (
    BibleCatalog,
    BibleEntry,
    CoverageResult,
    CoverageStatus,
    EntryCoverage,
)
from saarthi_ai.persistence import reporting_workflow
from saarthi_ai.persistence.database import SaarthiDatabase
from saarthi_ai.persistence.models import (
    AuditEventType,
    EvidenceCreate,
    EvidenceType,
    ExecutionCreate,
)
from saarthi_ai.persistence.reporting_workflow import run_tracked_report
from saarthi_ai.reporting.builder import build_report_model
from saarthi_ai.reporting.docx_renderer import render_report_docx
from saarthi_ai.reporting.json_renderer import render_report_json
from saarthi_ai.reporting.models import (
    EngagementMeta,
    ReportFinding,
    ReportModel,
    priority_band,
    severity_label,
)
from saarthi_ai.reporting.narrative import fallback_summary

ORCHESTRATION_ID = "orchestration-8a-test"


def _catalog() -> BibleCatalog:
    return BibleCatalog(
        entries=[
            BibleEntry.build(
                number=1,
                title="Cross-Site Scripting (XSS)",
                rating="High",
                description="Reflected XSS.",
                security_risk="Session theft.",
                recommendation="Encode output.",
                references="OWASP XSS",
                proof_of_concept="Step 1: inject.",
            ),
            BibleEntry.build(
                number=2,
                title="Cookie Flag HTTPOnly Not Set",
                rating="Low",
                description="Missing HttpOnly.",
                recommendation="Set HttpOnly.",
            ),
        ]
    )


def _coverage() -> CoverageResult:
    return CoverageResult(
        target="https://target.test",
        catalog_size=2,
        ai_used=True,
        coverage=[
            EntryCoverage(
                entry_id="cross-site-scripting-xss",
                title="Cross-Site Scripting (XSS)",
                severity=Severity.HIGH,
                status=CoverageStatus.CONFIRMED,
                source="ai",
            ),
            EntryCoverage(
                entry_id="cookie-flag-httponly-not-set",
                title="Cookie Flag HTTPOnly Not Set",
                severity=Severity.LOW,
                status=CoverageStatus.LIKELY,
                source="deterministic",
            ),
        ],
    )


@pytest.fixture
def database(tmp_path: Path) -> SaarthiDatabase:
    repository = SaarthiDatabase(tmp_path / "8a.db")
    repository.initialize()
    return repository


def _seed_run(database: SaarthiDatabase, tmp_path: Path) -> str:
    parent = database.create_execution(
        ExecutionCreate(
            assessment_name="Phase 8A",
            asset_types=["web"],
            targets=["https://target.test/"],
            authorization_confirmed=True,
            metadata={
                "orchestration_id": ORCHESTRATION_ID,
                "execution_role": "orchestration_parent",
            },
        )
    )
    body = json.dumps(_coverage().as_dict(), indent=2).encode("utf-8")
    path = tmp_path / "cov.json"
    path.write_bytes(body)
    database.add_evidence(
        parent.execution_id,
        EvidenceCreate(
            evidence_type=EvidenceType.BIBLE_COVERAGE_RESULT,
            source="bible-coverage",
            path=str(path),
            sha256="0" * 64,
            size_bytes=len(body),
            content_type="application/json",
        ),
    )
    return parent.execution_id


# --- model helpers -----------------------------------------------------------


def test_priority_band_and_labels() -> None:
    assert priority_band(Severity.CRITICAL)[0] == "P1"
    assert priority_band(Severity.LOW)[2] == "0.1-3.9"
    assert severity_label(Severity.INFO) == "Informational"


def test_report_model_summaries() -> None:
    model = ReportModel(
        engagement=EngagementMeta(),
        findings=[
            ReportFinding(
                finding_id="a", number=1, title="A", severity=Severity.LOW
            ),
            ReportFinding(
                finding_id="b", number=2, title="B", severity=Severity.CRITICAL
            ),
        ],
    )
    assert model.highest_severity is Severity.CRITICAL
    assert model.severity_counts["critical"] == 1
    # worst first
    assert model.sorted_findings[0].severity is Severity.CRITICAL


# --- builder -----------------------------------------------------------------


def test_build_report_model_maps_bible(
    database: SaarthiDatabase, tmp_path: Path
) -> None:
    _seed_run(database, tmp_path)
    model = build_report_model(
        database,
        orchestration_id=ORCHESTRATION_ID,
        engagement=EngagementMeta(company="Acme"),
        catalog=_catalog(),
    )
    assert len(model.findings) == 2
    xss = model.sorted_findings[0]
    assert xss.severity is Severity.HIGH
    assert xss.exploited_verified is True  # confirmed status
    assert xss.security_risk == "Session theft."
    assert xss.proof_of_concept.startswith("Step 1")
    assert model.overall_posture  # computed
    assert model.target == "https://target.test/"


# --- renderers ---------------------------------------------------------------


def test_render_report_json(tmp_path: Path) -> None:
    model = ReportModel(
        engagement=EngagementMeta(),
        findings=[
            ReportFinding(
                finding_id="a", number=1, title="A", severity=Severity.HIGH
            )
        ],
    )
    out = render_report_json(model, tmp_path / "r.json")
    payload = json.loads(out.read_text())
    assert payload["schema"] == "saarthi.report/2"
    assert payload["findings"][0]["priority"] == "P2"
    assert payload["severity_counts"]["high"] == 1
    # Extended schema surfaces roadmap + appendix data.
    assert "remediation_roadmap" in payload
    assert "methodology" in payload
    assert payload["severity_key"][0]["priority"] == "P1"


def test_render_report_docx_fallback(tmp_path: Path) -> None:
    model = ReportModel(
        engagement=EngagementMeta(app_name="https://target.test"),
        executive_summary="Summary text.",
        findings=[
            ReportFinding(
                finding_id="cross-site-scripting-xss",
                number=1,
                title="Cross-Site Scripting (XSS)",
                severity=Severity.HIGH,
                description="Reflected XSS.",
            )
        ],
    )
    out = render_report_docx(model, tmp_path / "r.docx", template_path=None)
    assert out.exists()
    document = Document(str(out))
    full = "\n".join(p.text for p in document.paragraphs)
    full += "\n" + "\n".join(
        c.text for t in document.tables for r in t.rows for c in r.cells
    )
    assert "Cross-Site Scripting (XSS)" in full
    assert "Summary text." in full
    assert "Reflected XSS." in full


def _make_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("Prepared for ORGNAME — APPNAME by TESTINGCOMPANY.")
    document.add_paragraph("Reviewed by REVIEWERNAME. Posture: OVERALLPOSTURE.")
    # scope table
    scope = document.add_table(rows=2, cols=2)
    scope.rows[0].cells[0].text = "S. No."
    scope.rows[0].cells[1].text = "URL(s)"
    scope.rows[1].cells[0].text = "1"
    # summary table
    summary = document.add_table(rows=2, cols=5)
    for cell, label in zip(
        summary.rows[0].cells,
        ("S. No.", "Title", "Risk", "Exploited / Verified", "ID"),
        strict=False,
    ):
        cell.text = label
    summary.rows[1].cells[0].text = "1"
    document.add_paragraph("[[FINDINGS HERE]]")
    document.save(str(path))


def test_render_report_docx_with_template(tmp_path: Path) -> None:
    template = tmp_path / "tpl.docx"
    _make_template(template)
    model = ReportModel(
        engagement=EngagementMeta(
            org_name="Demo Corp",
            app_name="https://target.test",
            company="Acme Security",
            reviewer="R. Lead",
        ),
        overall_posture="Needs Improvement",
        findings=[
            ReportFinding(
                finding_id="cross-site-scripting-xss",
                number=1,
                title="Cross-Site Scripting (XSS)",
                severity=Severity.HIGH,
                affected_urls=["https://target.test"],
                description="Reflected XSS.",
            )
        ],
    )
    out = render_report_docx(model, tmp_path / "out.docx", template_path=template)
    document = Document(str(out))
    full = "\n".join(p.text for p in document.paragraphs)
    full += "\n" + "\n".join(
        c.text for t in document.tables for r in t.rows for c in r.cells
    )
    assert "ORGNAME" not in full
    assert "TESTINGCOMPANY" not in full
    assert "[[FINDINGS HERE]]" not in full
    assert "Demo Corp" in full
    assert "Acme Security" in full
    assert "R. Lead" in full
    assert "Needs Improvement" in full
    assert "Cross-Site Scripting (XSS)" in full
    assert "Reflected XSS." in full


# --- narrative ---------------------------------------------------------------


def test_fallback_summary_variants() -> None:
    empty = ReportModel(engagement=EngagementMeta(app_name="app"))
    assert "No exploitable" in fallback_summary(empty)

    populated = ReportModel(
        engagement=EngagementMeta(app_name="app"),
        findings=[
            ReportFinding(
                finding_id="a", number=1, title="A", severity=Severity.HIGH
            )
        ],
    )
    assert "1 finding" in fallback_summary(populated)


# --- workflow ----------------------------------------------------------------


def test_run_tracked_report_persists(
    database: SaarthiDatabase, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    execution_id = _seed_run(database, tmp_path)
    monkeypatch.setattr(reporting_workflow, "report_template_path", lambda: None)
    monkeypatch.setattr(
        "saarthi_ai.reporting.builder.load_bible_catalog", _catalog
    )
    # Keep the test fast/deterministic regardless of a local LibreOffice.
    monkeypatch.setattr(
        reporting_workflow, "render_report_pdf", lambda *a, **k: None
    )

    tracked = run_tracked_report(
        database,
        execution_id,
        orchestration_id=ORCHESTRATION_ID,
        evidence_root=tmp_path / "reports",
        engagement=EngagementMeta(company="Acme"),
    )

    assert tracked.evidence is not None
    assert tracked.evidence.evidence_type is EvidenceType.ASSESSMENT_REPORT
    assert Path(tracked.docx_path).exists()
    assert Path(tracked.json_path).exists()
    assert Path(tracked.markdown_path).exists()
    assert Path(tracked.html_path).exists()
    assert tracked.pdf_path is None  # no converter in the test env
    assert "md" in tracked.evidence.metadata["formats"]
    assert tracked.evidence.metadata["findings"] == 2
    assert tracked.model.ai_narrative_used is False  # no summary passed

    events = database.list_audit_events(execution_id)
    assert any(
        e.event_type is AuditEventType.TOOL_COMPLETED and "[8A]" in e.message
        for e in events
    )
