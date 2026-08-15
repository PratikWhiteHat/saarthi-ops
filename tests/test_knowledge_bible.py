"""Vulnerability-library models, parser, and loader tests."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from saarthi_ai.exploit_confirmation.models import Severity
from saarthi_ai.knowledge.bible_parser import BibleParseError, parse_bible_docx
from saarthi_ai.knowledge.loader import (
    BibleNotAvailableError,
    bible_available,
    load_bible_catalog,
    save_bible_catalog,
)
from saarthi_ai.knowledge.models import (
    BibleCatalog,
    BibleEntry,
    derive_keywords,
    pick_severity,
    slugify,
)


def test_slugify_is_stable_and_safe() -> None:
    assert slugify("Cross-Site Scripting (XSS)") == "cross-site-scripting-xss"
    assert slugify("Cookie Flag “HTTPOnly” Not Set") == "cookie-flag-httponly-not-set"
    assert slugify("A & B") == "a-and-b"
    assert slugify("   ") == "entry"


def test_derive_keywords_drops_stopwords_and_short_tokens() -> None:
    keywords = derive_keywords("Account Lockout Not Implemented")
    assert "account" in keywords
    assert "lockout" in keywords
    assert "not" not in keywords  # stopword


def test_pick_severity_handles_compound_rating() -> None:
    assert pick_severity("Low/Medium (based on credentials header)") is Severity.MEDIUM
    assert pick_severity("High") is Severity.HIGH
    assert pick_severity("SEVERITY") is Severity.INFO


def test_bible_entry_build_derives_fields() -> None:
    entry = BibleEntry.build(
        number=3,
        title="Cross-Site Scripting (XSS)",
        rating="High",
        description="desc",
    )
    assert entry.entry_id == "cross-site-scripting-xss"
    assert entry.severity is Severity.HIGH
    assert entry.severity_rank == 3
    assert "scripting" in entry.keywords


def test_catalog_get() -> None:
    catalog = BibleCatalog(
        entries=[BibleEntry.build(number=1, title="Clickjacking", rating="Low")]
    )
    assert catalog.get("clickjacking") is not None
    assert catalog.get("missing") is None
    assert len(catalog) == 1


def test_save_and_load_round_trip(tmp_path: Path) -> None:
    catalog = BibleCatalog(
        version="1",
        source="unit",
        entries=[
            BibleEntry.build(number=1, title="HSTS Not Implemented", rating="Low"),
            BibleEntry.build(number=2, title="SQL Injection", rating="Critical"),
        ],
    )
    assert not bible_available(base=tmp_path)
    save_bible_catalog(catalog, base=tmp_path)
    assert bible_available(base=tmp_path)

    loaded = load_bible_catalog(base=tmp_path)
    assert len(loaded) == 2
    assert loaded.get("sql-injection").severity is Severity.CRITICAL


def test_load_missing_catalog_raises(tmp_path: Path) -> None:
    with pytest.raises(BibleNotAvailableError):
        load_bible_catalog(base=tmp_path)


def _make_library_docx(path: Path) -> None:
    document = Document()
    fields = [
        ("Vulnerability", "Clickjacking"),
        ("Affected URL(s)", "AFFECTEDURLS"),
        ("Vulnerability Rating", "Low"),
        ("Description", "Frames are allowed."),
        ("Security Risk", "UI redress."),
        ("Recommendation", "Set X-Frame-Options."),
        ("Reference(s)", "OWASP Clickjacking"),
        ("Proof of Concept (PoC)", "Step 1: frame it."),
    ]
    table = document.add_table(rows=len(fields), cols=2)
    for row, (label, value) in zip(table.rows, fields, strict=True):
        row.cells[0].text = label
        row.cells[1].text = value
    document.save(str(path))


def test_parse_bible_docx(tmp_path: Path) -> None:
    src = tmp_path / "lib.docx"
    _make_library_docx(src)
    catalog = parse_bible_docx(src)
    assert len(catalog) == 1
    entry = catalog.entries[0]
    assert entry.title == "Clickjacking"
    assert entry.severity is Severity.LOW
    assert entry.recommendation == "Set X-Frame-Options."
    assert entry.proof_of_concept.startswith("Step 1")


def test_parse_bible_docx_missing_file(tmp_path: Path) -> None:
    with pytest.raises(BibleParseError):
        parse_bible_docx(tmp_path / "nope.docx")


def test_parse_bible_docx_no_tables(tmp_path: Path) -> None:
    src = tmp_path / "empty.docx"
    Document().save(str(src))
    with pytest.raises(BibleParseError):
        parse_bible_docx(src)
